"""Document parsing: PDF (pdfplumber tables + layout-mode text, pypdf fallback), DOCX, Markdown/TXT."""
from __future__ import annotations

import html
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

from ..config import ApiError, MAX_UPLOAD_BYTES


ALLOWED_EXT = {".pdf", ".docx", ".md", ".txt"}


def safe_doc_name(name: str) -> str:
    base = Path(unicodedata.normalize("NFC", name).replace("\\", "/")).name.strip()
    base = re.sub(r"[^\w.\-&,() ]", "_", base)[:120]  # keeps "IAM, Endgeräte & Wechselmedien (v2).pdf" readable
    if not re.search(r"\.[A-Za-z0-9]{1,5}$", base):
        base += ".txt"
    return base or "dokument.txt"


MAX_UNZIPPED_BYTES = 64 * 1024 * 1024


def _check_zip_budget(data: bytes) -> None:
    """DOCX is a ZIP: reject archives whose declared uncompressed size is absurd *before* any parser allocates memory (zip bomb)."""
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile:
        raise ApiError(422, "DOCX_UNREADABLE", "DOCX konnte nicht gelesen werden: BadZipFile")
    if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
        raise ApiError(413, "DOCX_TOO_LARGE", "DOCX entpackt zu groß (möglicher Zip-Bomb).")


def parse_document(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ApiError(415, "UNSUPPORTED_TYPE", f"Dateityp {ext or '?'} nicht unterstützt (erlaubt: .pdf, .docx, .md, .txt).")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ApiError(413, "TOO_LARGE", f"Datei größer als {MAX_UPLOAD_BYTES // 1024 // 1024} MB.")
    if ext in (".md", ".txt"):
        for enc in ("utf-8-sig", "cp1252"):
            try:
                return _clean_text(data.decode(enc))
            except UnicodeDecodeError:
                continue
        return _clean_text(data.decode("utf-8", errors="replace"))
    if ext == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ApiError(422, "PDF_ENCRYPTED", "Verschlüsselte PDFs werden nicht unterstützt.")
            pages = reader.pages[:300]
            try:
                return "\n\n".join(_plumber_pages(data, pages))
            except Exception:  # noqa: BLE001 — pdfplumber failed on this file: pypdf layout heuristics for every page
                return "\n\n".join(_pdf_page_text(p) for p in pages)
        except ApiError:
            raise
        except Exception as e:  # noqa: BLE001
            raise ApiError(422, "PDF_UNREADABLE", f"PDF konnte nicht gelesen werden: {type(e).__name__}")
    _check_zip_budget(data)
    try:
        return _docx_to_text(data)
    except Exception as e:  # noqa: BLE001
        raise ApiError(422, "DOCX_UNREADABLE", f"DOCX konnte nicht gelesen werden: {type(e).__name__}")


# ── plain text / Markdown: export remnants and text tables ──
# Intranet/HTML exports leave entities and line-break tags ("&nbsp;", "&amp;", "<br>"); text tables come as Markdown pipe
# tables with blank lines between rows and wrapped cells, or as ASCII tables without outer pipes ("a | b" over "---+---").
# Both are normalised to one Markdown table so the chunker writes one "Spalte: Wert" statement per row.
HTML_BREAK = re.compile(r"<\s*(?:br|/p|p)\s*/?\s*>", re.I)
PIPE_ROW = re.compile(r"^\s*\|.*\|\s*$")
MD_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)*\|?\s*$")
ASCII_SEP = re.compile(r"^\s*[-=]{3,}(?:\s*[+|]\s*[-=]{3,})+\s*$")


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if re.search(r"&(?:[a-z]+|#\d+);|<\s*(?:br|p)\b", text, re.I):
        text = html.unescape(HTML_BREAK.sub("\n", text)).replace("\xa0", " ")
    return _text_tables("\n".join(_ascii_tables(text.split("\n"))))


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _ascii_tables(lines: list[str]) -> list[str]:
    """'Header | Header' over a '----+----' line and rows with the same number of ' | ' → Markdown pipe table."""
    out: list[str] = []
    i = 0
    while i < len(lines):
        head = lines[i]
        if i + 1 < len(lines) and " | " in head and not PIPE_ROW.match(head) and ASCII_SEP.match(lines[i + 1]):
            n = len(_cells(head))
            j = i + 2
            rows = []
            while j < len(lines) and " | " in lines[j] and len(_cells(lines[j])) == n:
                rows.append(_cells(lines[j]))
                j += 1
            if rows:
                out += ["| " + " | ".join(_cells(head)) + " |", "|" + "---|" * n] + ["| " + " | ".join(r) + " |" for r in rows]
                i = j
                continue
        out.append(head)
        i += 1
    return out


def _merge_cells(rows: list[list[str]]) -> list[str]:
    merged = [""] * max(len(r) for r in rows)
    for r in rows:
        for k, c in enumerate(r):
            merged[k] = _join_cell(merged[k], c) if c else merged[k]
    return merged


def _text_tables(text: str) -> str:
    """Markdown pipe tables: blank lines between rows are dropped; with blank lines the lines between two blank lines are
    one row (wrapped cells), without them a line with an empty first cell continues the row above."""
    lines = text.split("\n")
    out: list[str] = []
    i = 0
    while i < len(lines):
        if not (PIPE_ROW.match(lines[i]) and i + 1 < len(lines) and MD_SEP.match(lines[i + 1])):
            out.append(lines[i])
            i += 1
            continue
        n = len(_cells(lines[i]))
        groups: list[list[list[str]]] = [[]]
        j, gaps = i + 2, False
        while j < len(lines):
            if PIPE_ROW.match(lines[j]) and len(_cells(lines[j])) == n:
                groups[-1].append(_cells(lines[j]))
                j += 1
            elif not lines[j].strip():
                k = j
                while k < len(lines) and not lines[k].strip():
                    k += 1
                if k < len(lines) and PIPE_ROW.match(lines[k]) and len(_cells(lines[k])) == n and not MD_SEP.match(lines[k]):
                    gaps, j = True, k
                    if groups[-1]:
                        groups.append([])
                else:
                    break
            else:
                break
        groups = [g for g in groups if g]
        if gaps:
            body = [_merge_cells(g) for g in groups]
        else:
            body = []
            for r in (r for g in groups for r in g):
                if body and not r[0] and any(r):
                    body[-1] = _merge_cells([body[-1], r])
                else:
                    body.append(r)
        out += [lines[i], lines[i + 1]] + ["| " + " | ".join(r) + " |" for r in body]
        i = j
    return "\n".join(out)


PAGE_FOOTER = re.compile(r"\b(?:Seite|Page)\s+\d+\s+(?:von|of)\s+\d+\s*$", re.I)
LAYOUT_CELL = re.compile(r"\S+(?: {1,2}\S+)*")  # a cell = words separated by ≤ 2 spaces; columns are ≥ 3 spaces apart


def _layout_cells(line: str) -> list[tuple[int, str]]:
    return [(m.start(), m.group(0)) for m in LAYOUT_CELL.finditer(line)]


def _join_cell(a: str, b: str) -> str:
    if not a:
        return b
    return a + b if re.search(r"\w-$", a) and b[:1].isalnum() else f"{a} {b}"  # "ERP-" + "Auftragsbearbeitung"


LAYOUT_TOL = 8  # max. distance (characters) between a cell and its column start; centred cells drift a few characters


def _layout_rows(block: list[str], starts: list[int], one_value: bool = False) -> list[str] | None:
    """Cells → columns. None when any assignment is uncertain (cell far from every column start, reaching into the next
    column, or two cells of one line in the same column): the block is then treated as running text, never as a table
    with values in the wrong column. one_value: rows split without blank lines — a second line carrying a number into a
    cell that already has one is a merged row, not a wrapped cell (uncertain → None)."""
    cols = [""] * len(starts)
    for line in block:
        used: set[int] = set()
        for pos, text in _layout_cells(line):
            k = min(range(len(starts)), key=lambda j: (abs(starts[j] - pos), j))
            if abs(starts[k] - pos) > LAYOUT_TOL or k in used or (k + 1 < len(starts) and pos + len(text) > starts[k + 1] + 1):
                return None
            if one_value and re.search(r"\d", cols[k]) and re.search(r"\d", text):
                return None
            used.add(k)
            cols[k] = _join_cell(cols[k], text)
    return [c.replace("|", "\\|").strip() for c in cols]


CONT_END = re.compile(r"(?:[-/&,+(]|\b(?:or|and|und|oder|with|mit|of|der|des|für|for))$")


def _continues(prev: dict[int, str], cells: list[tuple[int, str]], starts: list[int]) -> bool:
    """Visible evidence that a line continues the cells above it: a cell above ends open ('ChaCha20-', 'Storage /', 'und',
    an unclosed bracket) or a cell starts lower case / with a closing bracket."""
    for pos, text in cells:
        k = min(range(len(starts)), key=lambda j: (abs(starts[j] - pos), j))
        above = prev.get(k, "")
        if text[:1].islower() or text[:1] == ")" or (above and (CONT_END.search(above) or above.count("(") > above.count(")"))):
            return True
    return False


def _row_groups(block: list[str], starts: list[int], header: bool) -> list[list[str]] | None:
    """Lines of a table without blank lines between rows → one line group per row, via the first column. A line without a
    first-column cell continues the current row (wrapped cells); a line with one starts a new row when it fills every
    column and shows no continuation evidence (_continues), continues the row when it shows such evidence, and is
    uncertain otherwise (partial line, capitalised: new row with an empty cell or a wrapped first cell?) → None, i.e.
    running text. Rows must also agree on numeric columns: a row without a number in a column where every other row has
    one is a wrapped line split off by mistake → None."""
    groups: list[list[str]] = []
    prev: dict[int, str] = {}
    for line in block:
        cells = _layout_cells(line)
        if not cells:
            continue
        first = abs(cells[0][0] - starts[0]) <= LAYOUT_TOL
        if groups and (not first or _continues(prev, cells, starts)):
            groups[-1].append(line)
        elif not groups or len(cells) == len(starts):
            groups.append([line])
            prev = {}
        else:
            return None
        for pos, text in cells:
            prev[min(range(len(starts)), key=lambda j: (abs(starts[j] - pos), j))] = text
    body = [_layout_rows(g, starts) for g in groups[1 if header else 0:]]
    if len(body) >= 2 and all(r is not None for r in body):
        for k in range(len(starts)):
            numeric = [bool(re.search(r"\d", r[k])) for r in body if r is not None and r[k]]
            if len(numeric) >= 3 and sum(numeric) == len(numeric) - 1:
                return None
    return groups


def _pdf_page_text(page: Any) -> str:
    """pypdf page → text via layout mode (_layout_md); plain extraction when the layout mode yields nothing."""
    try:
        raw = page.extract_text(extraction_mode="layout") or ""
    except Exception:  # noqa: BLE001 — older pypdf / odd fonts
        raw = ""
    if not raw.strip():
        return page.extract_text() or ""
    return _layout_md(raw)


def _layout_md(raw: str) -> str:
    """Layout-mode text keeps table columns: blocks with ≥ 3 aligned cells become Markdown tables (header + one row
    per blank-line separated block), two-line KPI boxes (values over UPPERCASE labels) become one-row tables, page footers
    are dropped."""
    blocks: list[list[str]] = []
    cur: list[str] = []
    for line in raw.split("\n"):
        line = line.rstrip()
        if PAGE_FOOTER.search(line):
            continue
        if line.strip():
            cur.append(line)
        elif cur:
            blocks.append(cur)
            cur = []
    if cur:
        blocks.append(cur)
    out: list[str] = []
    i = 0
    while i < len(blocks):
        head = blocks[i]
        widest = max(head, key=lambda ln: len(_layout_cells(ln)))
        if len(_layout_cells(widest)) >= 3:
            rows = []
            j = i + 1
            while j < len(blocks) and len(_layout_cells(blocks[j][0])) >= 2:
                rows.append(blocks[j])
                j += 1
            starts = [p for p, _ in _layout_cells(widest)]
            header, body = None, []
            if rows and len(head) <= 3:
                header = _layout_rows(head, starts)
                # rows separated by blank lines: one block = one row (wrapped cells may fill every column); a single
                # block after the header holds all rows without blank lines and is split by the first column
                groups = _row_groups(rows[0], starts, False) if len(rows) == 1 else rows
                tight = len(rows) == 1
                body = [_layout_rows(g, starts, tight) for g in groups] if groups else []
            elif not rows and len(head) >= 3:  # whole table in one block: header row + ≥ 2 rows split by the first column
                groups = _row_groups(head, starts, True)
                if groups and len(groups) >= 3:
                    header = _layout_rows(groups[0], starts)
                    body = [_layout_rows(g, starts, True) for g in groups[1:]]
            if header is not None and body and all(x is not None for x in body):
                out += ["", "| " + " | ".join(header) + " |", "|" + "---|" * len(starts)]
                out += ["| " + " | ".join(x) + " |" for x in body if x is not None] + [""]
                i = j
                continue
            if len(head) == 2:
                vals, labels = (_layout_cells(head[0]), _layout_cells(head[1]))
                letters = re.sub(r"[^A-Za-zÄÖÜäöüß]", "", head[1])
                centred = all(lp - 2 <= vp + len(vt) / 2 <= lp + len(lt) + 2 for (vp, vt), (lp, lt) in zip(vals, labels))
                if len(vals) == len(labels) and letters and letters.isupper() and centred:
                    out += ["", "| " + " | ".join(t for _, t in labels) + " |", "|" + "---|" * len(labels), "| " + " | ".join(t for _, t in vals) + " |", ""]
                    i += 1
                    continue
        out += [re.sub(r"\s{2,}", " ", ln.strip()) for ln in head] + [""]
        i += 1
    return "\n".join(out)


# ── pdfplumber tables (class b: ruled tables with wrapped cells and headers fall apart in the layout text) ──
# Per page: ruled tables (lines strategy). Cells are joined across line breaks, a header spread over two rows is merged, a
# table continued on the next page (same column positions, first on the page) keeps its header (a repeated header row is
# dropped). Text above/between/below tables goes through the layout heuristics as before; a page without ruled tables, or
# a pdfplumber error, uses the pypdf path unchanged. pdfplumber's text strategy is not used: on borderless tables it cut
# words across columns and assigned uncertain rows ("kau | m spuerbar"), where the layout heuristics keep running text.
LINES = {"vertical_strategy": "lines", "horizontal_strategy": "lines"}


def _cell(text: str | None) -> str:
    out = ""
    for part in (text or "").split("\n"):
        out = _join_cell(out, part.strip()) if part.strip() else out
    return out.replace("|", "\\|")


def _table_rows(raw: list[list[str | None]]) -> list[list[str]]:
    """Clean rows: a merged cell (None) inherits the value above (row span) or, in the first row, the value to its left
    (column span of a header); empty columns are dropped."""
    rows: list[list[str]] = []
    for r in raw:
        row = [_cell(c) if c is not None else (rows[-1][k] if rows and k < len(rows[-1]) else "") for k, c in enumerate(r)]
        if not rows:
            row = [c or (row[k - 1] if k and r[k] is None else c) for k, c in enumerate(row)]
        if any(row):
            rows.append(row)
    keep = [k for k in range(len(rows[0])) if any(r[k] for r in rows)] if rows else []
    return [[r[k] for k in keep] for r in rows]


def _header_rows(rows: list[list[str]], raw: list[list[str | None]]) -> int:
    """1, or 2 when the first row has merged/empty cells over a second row of labels without numbers (two-row header)."""
    if len(rows) >= 3 and (None in raw[0] or "" in raw[0]) and all(c and not re.search(r"\d", c) for c in rows[1]):
        return 2
    return 1


def _md_table(header: list[str], body: list[list[str]]) -> list[str]:
    header = [h or f"Spalte {k + 1}" for k, h in enumerate(header)]
    return ["", "| " + " | ".join(header) + " |", "|" + "---|" * len(header)] + ["| " + " | ".join(r) + " |" for r in body] + [""]


def _plumber_pages(data: bytes, pypdf_pages: list[Any]) -> list[str]:
    import pdfplumber

    out: list[str] = []
    prev: dict[str, Any] | None = None  # last table of the previous page: header + column x positions
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for n, page in enumerate(pdf.pages[: len(pypdf_pages)]):
            tables = []
            for t in sorted(page.find_tables(LINES), key=lambda t: t.bbox[1]):
                raw = t.extract()
                rows = _table_rows(raw)
                if len(rows) < 2 or len(rows[0]) < 2:
                    continue  # a single ruled box (callout) is running text
                cols = [round(c[0]) for c in t.rows[0].cells if c is not None]
                tables.append((t.bbox, rows, raw, cols))
            if not tables:
                runs = _column_runs(page)
                out.append(_columns_text(page, runs) if runs else _pdf_page_text(pypdf_pages[n]))
                prev = None
                continue
            parts: list[str] = []
            top = 0.0
            for k, (bbox, rows, raw, cols) in enumerate(tables):
                above = _band_text(page, top, bbox[1])
                if above.strip():
                    parts.append(above)
                first_on_page = k == 0 and len([ln for ln in above.splitlines() if ln.strip()]) <= 1
                if prev and first_on_page and len(rows[0]) == len(prev["header"]) and \
                        all(abs(a - b) <= 3 for a, b in zip(cols, prev["cols"])) and len(cols) == len(prev["cols"]):
                    header, body = prev["header"], rows[1:] if rows[0] == prev["header"] else rows  # continued table
                else:
                    h = _header_rows(rows, raw)
                    header = rows[0] if h == 1 else [_join_cell(a, b) if a != b else a for a, b in zip(rows[0], rows[1])]
                    body = rows[h:]
                if body:
                    parts += _md_table(header, body)
                prev = {"header": header, "cols": cols}
                top = bbox[3]
            below = _band_text(page, top, page.height)
            if below.strip():
                parts.append(below)
                prev = None  # running text after the table: the next page's table is a new one
            out.append("\n".join(parts))
    return out


def _lines(words: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    lines: list[list[dict[str, Any]]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(w["top"] - lines[-1][0]["top"]) <= 3:
            lines[-1].append(w)
        else:
            lines.append([w])
    return lines


def _column_runs(page: Any) -> list[tuple[float, float, float]]:
    """Two-column regions (class: two-column pages whose lines run into each other in the layout text) → [(top, bottom,
    gutter x)]: ≥ 6 consecutive lines that leave a common vertical gap ≥ 8 pt free in the middle 30–70 % of the page,
    each side carrying text in ≥ half of the lines with ≥ 3 words per line (prose; the columns' lines need not be at the
    same height). Borderless tables have short cells and stay with the layout path."""
    width = float(page.width)
    grid = [x for x in range(int(0.3 * width), int(0.7 * width), 2)]
    runs: list[tuple[float, float, float]] = []
    cur: list[list[dict[str, Any]]] = []
    free: set[int] = set()

    def close() -> None:
        if len(cur) < 6 or not free:
            return
        xs = sorted(free)
        best, start = (0, 0), xs[0]
        for a, b in zip(xs, xs[1:] + [None]):
            if b != a + 2:
                if a - start > best[1] - best[0]:
                    best = (start, a)
                start = b if b is not None else a
        if best[1] - best[0] < 8:
            return
        g = (best[0] + best[1]) / 2
        left = [sum(1 for w in ln if w["x1"] <= g) for ln in cur]
        right = [sum(1 for w in ln if w["x0"] >= g) for ln in cur]
        l_lines, r_lines = [a for a in left if a], [b for b in right if b]
        if len(l_lines) >= 0.5 * len(cur) and len(r_lines) >= 0.5 * len(cur) and \
                sum(l_lines) / len(l_lines) >= 3 and sum(r_lines) / len(r_lines) >= 3:
            runs.append((min(w["top"] for w in cur[0]) - 1, max(w["bottom"] for w in cur[-1]) + 1, g))

    last_free: set[int] = set()
    for ln in _lines(page.extract_words()):
        line_free = {x for x in grid if not any(w["x0"] - 4 < x < w["x1"] + 4 for w in ln)}
        if cur and free & line_free:
            cur.append(ln)
            free &= line_free
        else:
            close()
            # the previous line may open the next region (a column heading line right after full-width text)
            cur, free = ([cur[-1], ln], last_free & line_free) if cur and last_free & line_free else ([ln], line_free)
        last_free = line_free
    close()
    return runs


def _columns_text(page: Any, runs: list[tuple[float, float, float]]) -> str:
    """Page with two-column regions: full-width text above/between/below, each region left column then right column."""
    parts: list[str] = []
    top = 0.0
    for r_top, r_bottom, g in runs:
        for t in (_band_text(page, top, r_top), _band_text(page, r_top, r_bottom, 0, g), _band_text(page, r_top, r_bottom, g, None)):
            if t.strip():
                parts.append(t)
        top = r_bottom
    rest = _band_text(page, top, page.height)
    if rest.strip():
        parts.append(rest)
    return "\n\n".join(parts)


def _band_text(page: Any, top: float, bottom: float, left: float = 0.0, right: float | None = None) -> str:
    """Text of a horizontal band as layout text for _layout_md: words placed at column x / median character width (a gap
    below 0.6 × font height is one space), lines by word top (± 3 pt), a blank line where the gap exceeds 1.5 × the usual
    line pitch (paragraph break), as in pypdf's layout mode."""
    if bottom - top < 1:
        return ""
    band = page.crop((left, max(0.0, top), page.width if right is None else right, min(page.height, bottom)))
    words = band.extract_words()
    if not words:
        return ""
    widths = sorted((c["x1"] - c["x0"]) for c in page.chars if c.get("text", " ").strip()) or [5.0]
    xd = max(widths[len(widths) // 2], 1.0)
    lines = _lines(words)
    tops = [ln[0]["top"] for ln in lines]
    gaps = sorted(b - a for a, b in zip(tops, tops[1:]) if b - a > 0)
    pitch = gaps[len(gaps) // 2] if gaps else 12.0
    raw: list[str] = []
    for n, ln in enumerate(lines):
        if n and tops[n] - tops[n - 1] > 1.5 * pitch:
            raw.append("")
        text, last = "", None
        for w in sorted(ln, key=lambda w: w["x0"]):
            if last is not None and w["x0"] - last["x1"] < 0.6 * (w["bottom"] - w["top"]):
                text += " " + w["text"]  # word gap (larger fonts have wider spaces): one space, never a column gap
            else:
                col = round((w["x0"] - left) / xd)
                text = text.ljust(max(col, len(text) + 1) if text else col) + w["text"]
            last = w
        raw.append(text)
    return _layout_md("\n".join(raw))


def _docx_paragraph_line(p: Any) -> str:
    txt = p.text.strip()
    if not txt:
        return ""
    style = (p.style.name if p.style is not None else "") or ""
    if style == "Title":
        return f"# {txt}"
    if style.startswith("Heading"):
        return f"## {txt}"
    if "List" in style:
        return f"- {txt}"
    return txt


def _docx_table_lines(table: Any) -> list[str]:
    """DOCX table → Markdown table (first row = header) so the chunker can turn rows into self-contained statements."""
    rows: list[list[str]] = []
    for row in table.rows:
        seen: set[int] = set()
        cells: list[str] = []
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat the same underlying element
                continue
            seen.add(id(cell._tc))
            cells.append(re.sub(r"\s+", " ", cell.text).replace("|", "\\|").strip())
        rows.append(cells)
    if not rows:
        return []
    lines = ["", "| " + " | ".join(rows[0]) + " |"]
    if len(rows) > 1:
        lines.append("|" + "---|" * len(rows[0]))
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return lines + [""]


def _docx_to_text(data: bytes) -> str:
    """Paragraphs and tables in document order, so a table stays under the heading it belongs to."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            lines.append(_docx_paragraph_line(Paragraph(child, doc)))
        elif child.tag.endswith("}tbl"):
            lines.extend(_docx_table_lines(Table(child, doc)))
    return "\n".join(lines)
